import csv
from pathlib import Path
from urllib.parse import urljoin
import requests
from bs4 import BeautifulSoup
from .security import public_url


def read_document(file_path):
    path = Path(file_path)
    sections, notices, used = [], [], 0

    def add(text, location):
        nonlocal used
        remaining = 60000 - used
        if remaining <= 0:
            return
        text = text.strip()[:remaining]
        if text:
            sections.append({"text": text, "location": location})
            used += len(text)

    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".csv"}:
        add(path.read_text(encoding="utf-8-sig"), "Document")
    elif suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        for number, page in enumerate(reader.pages[:80], 1):
            add(page.extract_text() or "", f"Page {number}")
            if used >= 60000:
                break
        if len(reader.pages) > 80:
            notices.append("Only the first 80 PDF pages were processed.")
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            for sheet in book:
                rows = sheet.iter_rows(max_row=1000, max_col=30, values_only=True)
                add("\n".join(" | ".join(str(v) for v in row if v is not None) for row in rows), sheet.title)
                if used >= 60000:
                    break
            notices.append("Spreadsheet extraction is limited to 1,000 rows and 30 columns per sheet.")
        finally:
            book.close()
    elif suffix == ".pptx":
        from pptx import Presentation
        for number, slide in enumerate(Presentation(path).slides, 1):
            add("\n".join(s.text for s in slide.shapes if s.has_text_frame), f"Slide {number}")
            if used >= 60000:
                break
    else:
        raise ValueError("Unsupported file format.")
    if used >= 60000:
        notices.append("Document text was limited to 60,000 characters.")
    if not sections:
        raise ValueError(f"No readable text in {path.name}. Scanned PDFs need OCR first.")
    return {"title": path.name, "sections": sections, "notices": notices}


def fetch_source(url):
    with requests.Session() as session:
        session.trust_env = False
        for _ in range(5):
            url = public_url(url)
            with session.get(url, timeout=(3.0, 5.0), stream=True, allow_redirects=False,
                             headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36"}) as response:
                if response.is_redirect:
                    url = urljoin(url, response.headers["Location"])
                    continue
                response.raise_for_status()
                if "html" not in response.headers.get("Content-Type", "").lower():
                    raise ValueError("Only HTML websites are supported; upload PDF files separately.")
                data = bytearray()
                for part in response.iter_content(32768):
                    data.extend(part)
                    if len(data) > 3_000_000:
                        raise ValueError("Webpage exceeds the download limit.")
                break
        else:
            raise ValueError("Too many website redirects.")
    soup = BeautifulSoup(bytes(data), "html.parser")
    title = soup.title.get_text(" ", strip=True) if soup.title else url
    for tag in soup(["script", "style", "nav", "footer", "header", "aside", "noscript"]):
        tag.decompose()
    main = soup.find("main") or soup.find("article") or soup.body or soup
    text = " ".join(main.get_text(" ", strip=True).split())
    if len(text.split()) < 40:
        raise ValueError("This page has too little readable text.")
    return {"title": title[:300], "url": url, "source_role": "external_reference",
            "source_type": "web", "sections": [{"location": "Webpage", "text": text[:15000]}],
            "notices": ["Webpage text was shortened."] if len(text) > 15000 else []}


def discover_sources(request):
    candidates = [{"title": "Provided reference", "url": u, "snippet": "", "provided": True}
                  for u in request["source_urls"]]
    errors = []
    if request["use_web_search"]:
        try:
            from duckduckgo_search import DDGS
        except ImportError:
            try:
                from ddgs import DDGS
            except ImportError:
                DDGS = None
        if DDGS is not None:
            # If user already supplied primary reference URLs, keep search ultra-fast
            queries = [request["topic"]] if candidates else [request["topic"] + " " + (request.get("audience") or "")]
            for query in queries:
                try:
                    for item in DDGS(timeout=4).text(query, backend="lite", max_results=3):
                        url = item.get("href", "")
                        SKIP_DOMAINS = ("youtube.com", "youtu.be", "facebook.com", "instagram.com", "twitter.com", "x.com", "tiktok.com", "pinterest.com")
                        if url.startswith(("https://", "http://")) and not any(d in url.lower() for d in SKIP_DOMAINS):
                            candidates.append({"title": item.get("title", "")[:250], "url": url,
                                               "snippet": item.get("body", "")[:400], "provided": False})
                except Exception:
                    pass
        else:
            errors.append("Web search engine library not found. Running with direct references.")
    unique = {c["url"]: c for c in reversed(candidates)}
    ordered = list(reversed(list(unique.values())))[:12]
    for i, item in enumerate(ordered, 1):
        item["candidate_id"] = f"W{i}"
    return ordered, errors


def chunk_sources(sources, model):
    tokenizer = model.tokenizer
    budget = min(280, model.max_seq_length - 12)
    chunks = []
    for source in sources:
        number = 0
        for section in source["sections"]:
            text = section["text"]
            encoded = tokenizer(text, add_special_tokens=False, return_offsets_mapping=True,
                                truncation=False, verbose=False)
            offsets = encoded["offset_mapping"]
            for start in range(0, len(offsets), budget - 40):
                end = min(start + budget, len(offsets))
                passage = text[offsets[start][0]:offsets[end - 1][1]].strip()
                if not passage:
                    continue
                number += 1
                chunks.append({"source_id": source["source_id"],
                               "chunk_id": f"{source['source_id']}-C{number}",
                               "location": section["location"], "text": passage})
                if len(chunks) >= 25 or end == len(offsets):
                    break
        if len(chunks) >= 25:
            break
    return chunks
