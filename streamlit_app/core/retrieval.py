import csv
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse, urlunparse
import urllib.request
from bs4 import BeautifulSoup

try:
    from duckduckgo_search import DDGS
except ImportError:
    DDGS = None

try:
    from sentence_transformers import SentenceTransformer
    import torch
except ImportError:
    SentenceTransformer = None
    torch = None

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None


def clean_url(url: str) -> str:
    parsed = urlparse(url.strip())
    if parsed.scheme not in ("http", "https") or not parsed.hostname:
        raise ValueError("Source URLs must use HTTP or HTTPS.")
    return urlunparse(parsed._replace(fragment=""))


def fetch_web_page(url: str, timeout: int = 15) -> str:
    req = urllib.request.Request(
        url,
        headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"}
    )
    with urllib.request.urlopen(req, timeout=timeout) as response:
        content_type = response.headers.get("Content-Type", "").lower()
        if "text/html" not in content_type and "text/plain" not in content_type:
            raise ValueError(f"Skipping non-text content type: {content_type}")
        charset = response.headers.get_content_charset() or "utf-8"
        html = response.read().decode(charset, errors="replace")
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "noscript", "svg", "nav", "footer", "header"]):
        tag.decompose()
    return " ".join(soup.get_text(separator=" ").split())


def discover_sources(request: dict, max_results: int = 8) -> tuple[list[dict], list[str]]:
    candidates = [
        {"title": "User-supplied URL", "url": url, "snippet": "", "provided": True}
        for url in request.get("source_urls", [])
    ]
    errors = []
    if request.get("use_web_search", True) and DDGS:
        query = " ".join(filter(None, (request.get("topic", ""), request.get("audience", ""))))
        try:
            for item in DDGS(timeout=20).text(query, max_results=max_results):
                try:
                    url = clean_url(item.get("href", ""))
                except ValueError:
                    continue
                candidates.append({
                    "title": item.get("title", ""),
                    "url": url,
                    "snippet": item.get("body", "")[:600],
                    "provided": False,
                })
        except Exception as exc:
            errors.append(f"Web search discovery notice: {exc}")
    return candidates, errors


def read_uploaded_file(uploaded_file, role: str = "user_reference") -> dict:
    name = uploaded_file.name
    suffix = Path(name).suffix.lower()
    sections, notices = [], []

    try:
        if hasattr(uploaded_file, "seek"):
            uploaded_file.seek(0)
    except Exception:
        pass

    if suffix in (".txt", ".md"):
        content = uploaded_file.read().decode("utf-8", errors="replace")
        if content.strip():
            sections.append({"text": content.strip(), "location": "Full document"})
    elif suffix == ".pdf":
        if not PdfReader:
            raise ImportError("pypdf library is required for PDF parsing. Please install with: pip install pypdf")
        reader = PdfReader(uploaded_file)
        for num, page in enumerate(reader.pages, 1):
            text = (page.extract_text() or "").strip()
            if text:
                sections.append({"text": text, "location": f"Page {num}"})
    elif suffix == ".csv":
        content = uploaded_file.read().decode("utf-8", errors="replace").splitlines()
        rows = [f"Row {i}: {row}" for i, row in enumerate(content, 1) if row.strip()]
        if rows:
            sections.append({"text": "\n".join(rows), "location": "CSV rows"})
    elif suffix == ".xlsx":
        if not load_workbook:
            raise ImportError("openpyxl library is required for Excel parsing. Please install with: pip install openpyxl")
        wb = load_workbook(uploaded_file, read_only=True, data_only=False)
        for sheet in wb:
            rows = [" | ".join(str(c.value) for c in row if c.value is not None) for row in sheet.iter_rows()]
            sheet_text = "\n".join(filter(None, rows)).strip()
            if sheet_text:
                sections.append({"text": sheet_text, "location": f"Sheet: {sheet.title}"})
        wb.close()
    elif suffix == ".pptx":
        if not Presentation:
            raise ImportError("python-pptx library is required for PowerPoint parsing. Please install with: pip install python-pptx")
        prs = Presentation(uploaded_file)
        for num, slide in enumerate(prs.slides, 1):
            parts = [shape.text.strip() for shape in slide.shapes if shape.has_text_frame and shape.text.strip()]
            if parts:
                sections.append({"text": "\n".join(parts), "location": f"Slide {num}"})
    else:
        raise ValueError(f"Unsupported file format '{suffix}'. Supported formats: PDF, TXT, MD, CSV, XLSX, PPTX.")

    if not sections:
        raise ValueError(f"No extractable text found in '{name}'. Scanned PDFs or empty documents cannot be parsed without OCR.")

    total_chars = sum(len(s["text"]) for s in sections)
    total_words = sum(len(s["text"].split()) for s in sections)

    return {
        "title": name,
        "sections": sections,
        "notices": notices,
        "role": role,
        "char_count": total_chars,
        "word_count": total_words,
    }


def chunk_text(text: str, target_words: int = 150, overlap_words: int = 30) -> list[str]:
    words = text.split()
    if not words:
        return []
    chunks = []
    step = max(1, target_words - overlap_words)
    for i in range(0, len(words), step):
        chunk = " ".join(words[i:i + target_words])
        if len(chunk.split()) >= 20:
            chunks.append(chunk)
    return chunks or [" ".join(words)]


_EMBEDDER = None

def get_embedding_model(model_name: str = "intfloat/multilingual-e5-small"):
    global _EMBEDDER
    if _EMBEDDER is None and SentenceTransformer:
        try:
            device = "cuda" if torch and torch.cuda.is_available() else "cpu"
            _EMBEDDER = SentenceTransformer(model_name, device=device)
        except Exception:
            _EMBEDDER = None
    return _EMBEDDER


def search_evidence(query: str, all_chunks: list[dict], top_k: int = 4) -> list[dict]:
    if not all_chunks:
        return []

    embedder = get_embedding_model()
    if embedder is not None:
        texts = [c["text"] for c in all_chunks]
        doc_embeddings = embedder.encode(["passage: " + t for t in texts], normalize_embeddings=True)
        query_embedding = embedder.encode(["query: " + query], normalize_embeddings=True)
        scores = (doc_embeddings @ query_embedding.T).flatten()
        top_indices = scores.argsort()[::-1][:top_k]
        return [all_chunks[idx] for idx in top_indices]
    else:
        # Fallback keyword overlap ranking
        query_terms = set(re.findall(r"\w+", query.lower()))
        ranked = []
        for c in all_chunks:
            chunk_terms = set(re.findall(r"\w+", c["text"].lower()))
            overlap = len(query_terms & chunk_terms)
            ranked.append((overlap, c))
        ranked.sort(key=lambda x: x[0], reverse=True)
        return [item[1] for item in ranked[:top_k]]
